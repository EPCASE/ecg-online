# -*- coding: utf-8 -*-
"""Remplit le dossier de candidature Bourse CNEC 2026 pour EDU-ECG.
Entrée : C:\\Users\\Administrateur\\Downloads\\PB2026-Bourse-CNEC-Candidature.docx
Sortie : C:\\Users\\Administrateur\\Desktop\\ECG lecture\\CNEC\\PB2026-Bourse-CNEC-Candidature_EDU-ECG.docx
Les champs personnels non connus sont marqués [À COMPLÉTER]."""
import copy
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

SRC = r"C:\Users\Administrateur\Downloads\PB2026-Bourse-CNEC-Candidature.docx"
DST = r"C:\Users\Administrateur\Desktop\ECG lecture\CNEC\PB2026-Bourse-CNEC-Candidature_EDU-ECG.docx"

d = docx.Document(SRC)
TODO = "[À COMPLÉTER]"


def find(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


def insert_after(par, text, bold=False, italic=False, size=10.5, bullet=False, align=None):
    new_p = copy.deepcopy(par._p)
    par._p.addnext(new_p)
    np_ = Paragraph(new_p, par._parent)
    for r in list(np_.runs):
        r._r.getparent().remove(r._r)
    np_.style = d.styles["Normal"]
    pf = np_.paragraph_format
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)
    pf.left_indent = Pt(12) if bullet else Pt(0)
    pf.first_line_indent = None
    if align == "j":
        np_.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = np_.add_run(("• " if bullet else "") + text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    return np_


def fill_section(title_prefix, blocks):
    cur = find(title_prefix)
    for kind, text in blocks:
        if kind == "h":
            cur = insert_after(cur, text, bold=True, size=11)
        elif kind == "b":
            cur = insert_after(cur, text, bullet=True, align="j")
        elif kind == "i":
            cur = insert_after(cur, text, italic=True, size=10)
        else:
            cur = insert_after(cur, text, align="j")


def set_field(prefix, value):
    find(prefix).add_run(" " + value)


# ── Identité / situation ────────────────────────────────────────────────────
for lab in ["Prénom", "Nom", "Date de naissance", "Adresse postale personnelle",
            "E-mail personnel", "Poste occupé", "Établissement professionnel",
            "Adresse postale professionnelle", "Tel.", "E-mail professionnel",
            "N° membre de la SFC"]:
    set_field(lab, TODO)
set_field("Date du dépôt de la demande", "octobre 2026")

insert_after(find("TITRE DU PROJET"),
    "EDU-ECG — Plateforme ouverte d'entraînement à l'interprétation de l'ECG en texte libre, "
    "avec correction neurosymbolique explicable et parcours adaptatifs pour le DES de médecine "
    "cardiovasculaire et le 2e cycle", bold=True, size=11)

# ── CV ──────────────────────────────────────────────────────────────────────
fill_section("CURRICULUM VITAE", [
    ("i", f"{TODO} — CV en une page maximum : fonctions actuelles, formation, responsabilités "
          "pédagogiques (enseignement ECG 2e/3e cycle), publications ou communications en pédagogie "
          "médicale, compétences en ingénierie pédagogique numérique."),
])

# ── Co-porteurs ─────────────────────────────────────────────────────────────
fill_section("CO-PORTEURS ET AUTRES", [
    ("p", f"Co-porteur(s) : {TODO} (enseignant(s) de cardiologie associé(s) à la relecture des cas et à la validation des barèmes)."),
    ("p", f"Établissement porteur / organisme gestionnaire : {TODO} (Université / CHU / association loi 1901)."),
    ("p", "Établissements partenaires sollicités pour la phase multicentrique : deux à trois services de "
          "cardiologie universitaires volontaires (lettres d'intention en cours), afin de tester l'ouverture "
          "de la plateforme à des enseignants extérieurs et de constituer la cohorte de l'étude d'impact."),
])

# ── Résumé (15 lignes) ──────────────────────────────────────────────────────
fill_section("RESUMÉ DU PROJET", [
    ("p", "L'interprétation de l'ECG est une compétence fondamentale et pourtant mal maîtrisée en fin de "
          "2e cycle comme en début de DES : l'entraînement repose sur l'exposition opportuniste et sur des "
          "QCM qui n'évaluent ni la rédaction d'une interprétation complète ni la hiérarchisation des "
          "anomalies. EDU-ECG est une plateforme web déjà opérationnelle dans laquelle l'étudiant rédige son "
          "interprétation en texte libre, comme au lit du malade, et reçoit immédiatement un score sur 100, "
          "la liste des éléments trouvés, manqués et erronés, et un commentaire pédagogique. Son originalité "
          "est un moteur de correction neurosymbolique : l'IA générative n'attribue pas la note, elle extrait "
          "les concepts du texte, qui sont ensuite scorés de façon déterministe contre une ontologie ECG de "
          "358 concepts (hiérarchie, prérequis, incompatibilités). Le score est donc reproductible, explicable "
          "et réglable par l'enseignant. Une calibration sur 508 copies réelles, dont 139 annotées en aveugle "
          "par un expert, a montré une corrélation machine/expert de 0,70 (Spearman). La bourse CNEC financera "
          "l'étape suivante : (1) un curriculum complet de 15 parcours couvrant les 75 cas avec réactivation "
          "espacée adaptative ; (2) l'ouverture multicentrique (comptes enseignants, import de tracés, barèmes "
          "locaux) ; (3) une étude contrôlée d'impact sur l'apprentissage et la rétention. Le code est ouvert "
          "et le coût marginal par étudiant quasi nul, garantissant la pérennité et le partage à l'échelle "
          "nationale du DES de médecine cardiovasculaire."),
])

# ── État des lieux (1 page) ─────────────────────────────────────────────────
fill_section("ETAT DES LIEUX ET ENJEUX", [
    ("h", "1. Une compétence critique, un apprentissage peu structuré"),
    ("p", "L'ECG est l'examen complémentaire le plus prescrit en médecine et le premier geste diagnostique "
          "de toute urgence cardiovasculaire. Sa lecture est exigible dès le 2e cycle (items EDN) et constitue "
          "le socle du DES de médecine cardiovasculaire. Pourtant, la littérature internationale montre de "
          "façon constante des taux d'erreur élevés chez les internes et les jeunes médecins, y compris sur "
          "des diagnostics à risque vital (BAV de haut grade, tachycardies à QRS larges, hyperkaliémie). La "
          "cause est connue : la lecture de l'ECG est une compétence de reconnaissance de formes qui exige un "
          "volume important de cas variés, une pratique délibérée et un feedback immédiat et spécifique. Or "
          "l'exposition en stage est opportuniste, inégale selon les terrains, et le retour de l'enseignant "
          "est rare, différé et non standardisé."),
    ("h", "2. Les outils existants n'évaluent pas la bonne compétence"),
    ("p", "Les banques d'entraînement disponibles reposent très majoritairement sur des QCM ou des réponses "
          "à choix fermés. Elles mesurent la reconnaissance d'une étiquette, pas la compétence réellement "
          "attendue : produire une interprétation complète, hiérarchisée et sûre (rythme, fréquence, "
          "conduction, axe, hypertrophies, repolarisation, conclusion). Elles ne détectent pas les erreurs "
          "dangereuses (affirmer un diagnostic incompatible), ne distinguent pas une hypothèse évoquée puis "
          "écartée d'une affirmation, et ne savent pas expliquer la note. Les correcteurs entièrement fondés "
          "sur un grand modèle de langage sont, eux, peu reproductibles, opaques et difficilement réglables "
          "par l'enseignant."),
    ("h", "3. Ce qui existe déjà : EDU-ECG, un prototype validé"),
    ("p", "EDU-ECG (ECG Online) est une application web développée au sein de l'équipe pédagogique et déjà "
          "utilisée par des étudiants. Elle comprend 75 cas ECG curés (tracé, contexte clinique, "
          "interprétation de référence rédigée par l'enseignant, rattachement aux items EDN). L'étudiant "
          "rédige son interprétation en texte libre et obtient en quelques secondes un score sur 100, le "
          "détail des éléments trouvés, manqués et erronés, et un commentaire pédagogique sourcé. Le moteur "
          "est neurosymbolique : extraction des concepts par IA (NER + juge), puis scoring déterministe sur "
          "une ontologie ECG de 358 concepts intégrant hiérarchie, prérequis, alternatives et "
          "incompatibilités. Le score combine une dimension d'adéquation (complétude) et une dimension de "
          "sécurité clinique (absence d'affirmation dangereuse). Une interface de curation permet à "
          "l'enseignant de régler le barème de chaque cas sans toucher au code."),
    ("p", "La démarche d'évaluation a été rigoureuse : 508 copies réelles ont été rejouées et 139 notées en "
          "aveugle par un cardiologue enseignant. Après calibration itérative (arbitrages cliniques tracés et "
          "documentés), la corrélation machine/expert atteint 0,70 (Spearman ; Pearson 0,75 ; erreur absolue "
          "moyenne 17 points), avec un biais légèrement sévère de la machine (−9 points) — un ordre de "
          "grandeur comparable à l'accord inter-correcteurs humain sur des copies rédactionnelles."),
    ("h", "4. Les enjeux : passer du prototype à l'outil partagé"),
    ("b", "Pédagogique : seuls 26 des 75 cas sont scénarisés dans cinq micro-parcours ; 49 ne sont "
          "accessibles qu'en banque libre. L'exposition est inégale et il n'existe ni progression explicite "
          "ni réactivation espacée, alors que la rétention est le véritable enjeu."),
    ("b", "Institutionnel : la plateforme est monocentrique. Pour servir le DES à l'échelle nationale, tout "
          "enseignant doit pouvoir y déposer ses tracés, régler ses barèmes et suivre ses étudiants sans "
          "dépendre de l'équipe fondatrice."),
    ("b", "Scientifique : l'impact sur l'apprentissage réel (et non sur la seule satisfaction) n'a pas encore "
          "été mesuré de façon contrôlée — condition d'une diffusion crédible."),
    ("b", "Soutenabilité : l'architecture (code ouvert, pipeline autonome, coût d'inférence de quelques "
          "centimes par copie) rend le passage à l'échelle réaliste moyennant un investissement ponctuel de "
          "développement — exactement l'objet de cet appel."),
])

# ── Descriptif (1,5 page) ───────────────────────────────────────────────────
fill_section("DESCRIPTIF DU PROJET", [
    ("p", "Le projet comporte trois axes de développement nouveaux, conduits sur l'année universitaire "
          "2026-2027, et s'appuie sur le socle technique existant (moteur de correction, ontologie, banque de "
          "75 cas), dont le fonctionnement courant n'est pas financé par la bourse."),
    ("h", "Axe 1 — Curriculum complet et apprentissage adaptatif"),
    ("b", "Construire 15 parcours fondamentaux de 5 ECG couvrant exactement les 75 cas, sans répétition fixe, "
          "organisés par famille sémiologique (troubles de conduction, arythmies atriales, tachycardies à QRS "
          "fins / larges, ischémie, hypertrophies, canalopathies, troubles ioniques, pièges techniques). "
          "Chaque parcours suit une progression en cinq temps : fondation (cas prototypique guidé), guidage "
          "(indices dégressifs), contraste (cas voisin à discriminer), intégration (cas complexe) et maîtrise "
          "(sans aide)."),
    ("b", "Implémenter un moteur de réactivation espacée : le choix du cas suivant dépend de l'historique "
          "individuel (concepts échoués, délai depuis la dernière réussite, confiance déclarée), selon un "
          "algorithme de répétition espacée appliqué aux concepts de l'ontologie plutôt qu'aux cas."),
    ("b", "Séparer explicitement trois moteurs : mesure (déterministe, auditable, existant), narratif "
          "(reformulation du résultat en feedback sur la tâche, le raisonnement et l'autorégulation) et "
          "adaptation (choix du cas et du niveau d'aide)."),
    ("b", "Tableau de bord étudiant distinguant couverture (cas rencontrés), maîtrise (concepts validés sans "
          "aide), rétention (concepts revalidés après espacement) et calibration (écart entre confiance "
          "déclarée et performance)."),
    ("h", "Axe 2 — Ouverture multicentrique : la plateforme comme bien commun du DES"),
    ("b", "Espace enseignant : comptes par service, import de tracés propres (image anonymisée + "
          "interprétation de référence), rattachement à l'ontologie assisté par IA avec validation humaine, "
          "réglage local des barèmes via l'interface de curation existante."),
    ("b", "Espace promotion : groupes d'étudiants, attribution de parcours, suivi agrégé (concepts les moins "
          "maîtrisés, copies à relire en priorité), export CSV pour les commissions pédagogiques."),
    ("b", "Gouvernance éditoriale : un comité de relecture de 3 à 5 enseignants de cardiologie (dont les "
          "co-porteurs) valide les cas partagés dans la banque commune ; les cas locaux restent privés par "
          "défaut. Chaque arbitrage de barème est tracé et versionné, comme c'est déjà le cas dans la "
          "documentation du projet."),
    ("b", "Accessibilité : application web sans installation, gratuite pour étudiants et enseignants, code "
          "source ouvert, conformité RGPD (données pseudonymisées, hébergement UE)."),
    ("h", "Axe 3 — Étude d'impact pédagogique"),
    ("b", "Design : étude prospective contrôlée sur deux à trois sites, externes de 5e-6e année et internes "
          "de phase socle. Groupe intervention (parcours adaptatifs EDU-ECG, 8 semaines) vs groupe contrôle "
          "(banque libre, même contenu sans scénarisation ni espacement). Critère principal : score à un test "
          "d'interprétation en texte libre de 10 ECG inédits, corrigé par deux experts en aveugle. Critères "
          "secondaires : rétention à 3 mois, taux d'erreurs « dangereuses », calibration de la confiance, "
          "temps d'usage, acceptabilité (questionnaire validé)."),
    ("b", "Validation du correcteur : extension de la calibration machine/expert à 300 copies annotées par "
          "deux correcteurs indépendants ; publication des métriques d'accord (Spearman, CCC de Lin, "
          "Bland-Altman) et de la matrice d'erreurs par famille de diagnostic."),
    ("b", "Valorisation : communication aux JESFC 2028 et au congrès de pédagogie médicale (SIFEM), article "
          "dans une revue d'éducation médicale, mise à disposition de la banque validée et du code sous "
          "licence ouverte avec mention du soutien du CNEC."),
    ("h", "Organisation et rôles"),
    ("p", "Le porteur assure la direction pédagogique et scientifique, la curation des cas et les arbitrages "
          "cliniques. Un développeur / ingénieur pédagogique, financé ponctuellement par la bourse, réalise "
          "les axes 1 et 2 sous spécification du porteur. Les co-porteurs animent le comité de relecture et "
          "coordonnent les sites de l'étude. Un point d'avancement mensuel et un rapport d'activité à 12 mois "
          "sont prévus, conformément à l'article 2 du règlement."),
])

# ── Périmètre (1 page) ──────────────────────────────────────────────────────
fill_section("PERIMETRE ET PUBLIC", [
    ("h", "Publics visés"),
    ("b", "2e cycle (DFASM) : étudiants de 4e à 6e année préparant les EDN — les 75 cas sont rattachés aux "
          "items du programme ; les parcours « fondation » et « guidage » leur sont destinés en priorité. "
          "Volume cible année 1 : promotions des sites partenaires (300 à 600 étudiants)."),
    ("b", "3e cycle : internes du DES de médecine cardiovasculaire (phase socle et approfondissement) et, par "
          "extension, internes de médecine d'urgence, anesthésie-réanimation, médecine interne et gériatrie. "
          "Les parcours « contraste », « intégration », « maîtrise » et la dimension sécurité clinique les "
          "concernent directement."),
    ("b", "Formation continue : médecins généralistes et urgentistes souhaitant entretenir leur compétence ; "
          "le moteur de réactivation espacée est conçu pour l'entretien à long terme. Ouverture en année 2, "
          "après validation de l'étude d'impact."),
    ("b", "Enseignants de cardiologie : public « producteur » — tout enseignant membre du CNEC pourra créer "
          "son espace, déposer ses tracés et suivre ses étudiants."),
    ("h", "Périmètre de contenu"),
    ("p", "La banque initiale de 75 cas couvre la sémiologie ECG exigible : rythme sinusal et variantes, "
          "troubles de conduction auriculo-ventriculaire et intraventriculaire, arythmies atriales et "
          "ventriculaires, tachycardies à QRS fins et larges, syndromes coronariens, hypertrophies, "
          "péricardite, canalopathies, troubles ioniques, stimulation cardiaque, pièges techniques (inversion "
          "d'électrodes, artéfacts). L'ouverture multicentrique permettra son extension progressive par les "
          "enseignants partenaires, sous contrôle du comité de relecture."),
    ("h", "Périmètre géographique et modalités d'accès"),
    ("p", "Année 1 (2026-2027) : site fondateur + deux à trois sites universitaires partenaires pour l'étude "
          "d'impact. Année 2 : ouverture à l'ensemble des enseignants du CNEC et diffusion via la SFC. Accès "
          "web (ordinateur, tablette, mobile), sans installation, gratuit, en français. L'outil est conçu "
          "pour s'intégrer aux enseignements existants (séances d'ECG dirigées, auto-entraînement, "
          "remédiation ciblée après évaluation) et non pour les remplacer."),
    ("h", "Ce que le projet n'est pas"),
    ("p", "EDU-ECG n'est pas un outil d'aide au diagnostic en soins ; il n'a aucune vocation clinique et "
          "n'est utilisé que sur des tracés pédagogiques anonymisés. Il ne se substitue pas au jugement de "
          "l'enseignant : chaque barème reste sous son contrôle et chaque score est explicable concept par "
          "concept."),
])

# ── Indicateurs (0,5 page) ──────────────────────────────────────────────────
fill_section("INDICATEURS DE SUIVI", [
    ("h", "Réalisation (suivi à 6 et 12 mois)"),
    ("b", "Couverture curriculaire : cas scénarisés dans un parcours (cible 75/75, départ 26/75)."),
    ("b", "Moteur adaptatif en production ; part des sessions utilisant la réactivation espacée."),
    ("b", "Espaces enseignants créés et sites actifs (cible ≥ 3) ; cas importés par des enseignants extérieurs (cible ≥ 25)."),
    ("b", "Inclusions dans l'étude d'impact (cible ≥ 150 étudiants) ; taux de complétion du protocole."),
    ("h", "Qualité du correcteur"),
    ("b", "Accord machine/expert en aveugle : Spearman ≥ 0,75 et CCC de Lin ≥ 0,75 sur 300 copies (départ 0,70 / 0,73 sur 139)."),
    ("b", "Faux positifs « erreur dangereuse » < 5 % ; scores contestés par les enseignants après relecture < 5 %."),
    ("h", "Impact pédagogique (critères de l'étude)"),
    ("b", "Différence de score au test final d'interprétation libre entre groupes (critère principal)."),
    ("b", "Rétention à 3 mois ; évolution du taux d'erreurs dangereuses ; calibration de la confiance."),
    ("b", "Usage : étudiants actifs, copies corrigées, temps médian par cas ; acceptabilité (SUS ≥ 70)."),
    ("h", "Diffusion"),
    ("b", "Communications (JESFC, SIFEM), article soumis, établissements utilisateurs à 24 mois, dépôt du code et de la banque sous licence ouverte."),
])

# ── Agenda (0,5 page) ───────────────────────────────────────────────────────
fill_section("AGENDA ET RETRO PLANNING", [
    ("b", "Nov.–déc. 2026 : spécifications des 15 parcours ; constitution du comité de relecture ; conventions avec les sites partenaires ; dépôt du protocole de l'étude (avis éthique, déclaration RGPD)."),
    ("b", "Janv. 2027 : remise de la bourse aux JESFC ; recrutement du développeur / ingénieur pédagogique."),
    ("b", "Janv.–mars 2027 : Axe 1 — moteur de parcours et tableau de bord ; scénarisation des 75 cas ; tests internes."),
    ("b", "Avr.–juin 2027 : Axe 2 — espace enseignant, import de cas, espace promotion, exports ; onboarding des sites partenaires ; extension de la calibration à 300 copies."),
    ("b", "Sept.–nov. 2027 : Axe 3 — étude d'impact (8 semaines d'intervention + test final) sur les promotions de rentrée."),
    ("b", "Déc. 2027 : rapport d'activité à 12 mois pour le jury (versement du solde) ; analyse intermédiaire."),
    ("b", "Janv.–mars 2028 : test de rétention à 3 mois ; analyse finale ; communication aux JESFC 2028 ; ouverture à l'ensemble des enseignants du CNEC ; soumission de l'article."),
])

# ── Budget ──────────────────────────────────────────────────────────────────
t = d.tables[1]


def setc(r, c, txt):
    cell = t.cell(r, c)
    cell.text = txt
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = Pt(9.5)


rows = [
    ("Prestation de développement / ingénierie pédagogique (axes 1 et 2), ~65 jours", "", "", "10 000"),
    ("Hébergement sécurisé (UE), nom de domaine, sauvegardes — 24 mois (autofinancé)", "", "—", ""),
    ("Coût d'inférence IA (extraction / feedback) — autofinancé", "", "—", ""),
    ("", "", "", ""),
]
for i, (lab, e, f, p_) in enumerate(rows):
    setc(2 + i, 0, lab); setc(2 + i, 1, e); setc(2 + i, 2, f); setc(2 + i, 3, p_)

setc(6, 0, "Montant global de la demande à l'AAP CNEC 2026")
setc(6, 1, "0"); setc(6, 2, "0"); setc(6, 3, "10 000")

setc(8, 0, "Autofinancement établissement : temps enseignant du porteur et des co-porteurs (valorisé)")
setc(8, 3, "[À COMPLÉTER — valorisation indicative]")
setc(9, 0, "Autofinancement porteur : hébergement, inférence IA, annotation, communication (~3 000 €)")
setc(9, 2, "3 000")
setc(10, 0, "Existant non refinancé : plateforme, moteur, ontologie, 75 cas (déjà développés)")
setc(10, 1, "—"); setc(10, 2, "—"); setc(10, 3, "—")
setc(11, 0, "Autre demande de financement en cours")
setc(11, 2, "[À COMPLÉTER ou « aucune »]")

setc(12, 0, "Montant du projet (hors valorisation du temps enseignant)")
setc(12, 1, "0"); setc(12, 2, "3 000"); setc(12, 3, "10 000")
setc(13, 0, "Part de financement demandé (%)")
setc(13, 2, "77 % (10 000 € sur 13 000 €)")

fill_section("BUDGET PRÉVISIONNEL", [
    ("i", "Total demandé : 10 000 €, intégralement affecté au poste personnel : financement ponctuel d'un "
          "développeur / ingénieur pédagogique (~65 jours) pour les fonctionnalités nouvelles — parcours "
          "adaptatifs et réactivation espacée (axe 1), espace enseignant multicentrique (axe 2). Conformément "
          "à l'article 1, la bourse ne finance aucune action récurrente existante : le socle technique "
          "(moteur, ontologie, banque de cas) est déjà développé, et les frais de fonctionnement "
          "(hébergement, inférence IA, annotation, communication, ~3 000 €) sont autofinancés par le porteur."),
])

# ── Devis / RIB ─────────────────────────────────────────────────────────────
fill_section("JUSTIFICATION DES DEMANDES", [
    ("i", f"{TODO} — joindre le devis du prestataire de développement / ingénierie pédagogique "
          "(~65 jours, 10 000 € TTC), détaillant les lots : moteur de parcours et réactivation espacée, "
          "tableau de bord, espace enseignant (comptes, import de cas, curation), espace promotion et exports."),
])
fill_section("COORDONNÉES BANCAIRES", [
    ("i", f"{TODO} — RIB de l'organisme gestionnaire (Université / CHU / association loi 1901) et de la "
          "personne référente, à renseigner dans les champs ci-dessous."),
])

d.save(DST)
print("OK ->", DST)
